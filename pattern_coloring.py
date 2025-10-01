"""Utility for parsing poker hand patterns and exporting them to Excel.

This module reads a CSV file that contains two columns:

* ``pattern`` – a textual description of cards, where parentheses indicate
  cards that share a suit.
* ``combinations`` – the number of distinct combinations available for the
  pattern.

The script parses each pattern, expands it into individual card positions, and
exports an Excel workbook where every card appears in its own column. Cards
that share a suit are highlighted with the same fill colour, making it easy to
spot suited groupings at a glance.

Example
-------

Running the script from the command line::

    python pattern_coloring.py patterns.csv patterns.xlsx

See the module-level ``main`` function for additional CLI options.
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill


@dataclass
class ParsedPattern:
    """Represents a fully parsed pattern.

    Attributes
    ----------
    pattern: str
        The raw textual pattern from the CSV file.
    combinations: str
        Combination count as a string. Left as ``str`` to preserve the raw
        value in case the CSV stores non-integer data.
    suit_groups: List[List[str]]
        Nested list representing cards that share a suit. Each inner list holds
        the card symbols that belong to the same suit.
    """

    pattern: str
    combinations: str
    suit_groups: List[List[str]]

    @property
    def flattened_cards(self) -> List[Tuple[str, int]]:
        """Return cards paired with their suit group index.

        The result is suitable for iterating in order when populating the Excel
        worksheet.
        """

        cards: List[Tuple[str, int]] = []
        for suit_index, cards_in_suit in enumerate(self.suit_groups):
            cards.extend((card, suit_index) for card in cards_in_suit)
        return cards


def parse_pattern(pattern: str) -> List[List[str]]:
    """Parse a pattern string into suit groupings.

    Parentheses indicate that the enclosed cards share a suit. Characters that
    are not inside parentheses are treated as individual cards, each belonging
    to its own suit.

    Parameters
    ----------
    pattern:
        Pattern string like ``"(AK)(A3)2"``.

    Returns
    -------
    List[List[str]]
        Nested list where each inner list contains the cards for one suit.
    """

    suits: List[List[str]] = []
    i = 0
    while i < len(pattern):
        char = pattern[i]
        if char == "(":
            try:
                closing = pattern.index(")", i + 1)
            except ValueError as exc:  # pragma: no cover - defensive programming
                raise ValueError(f"Unmatched '(' in pattern: {pattern}") from exc

            group = list(pattern[i + 1 : closing])
            if not group:
                raise ValueError(f"Empty parentheses found in pattern: {pattern}")
            suits.append(group)
            i = closing + 1
        else:
            j = i
            while j < len(pattern) and pattern[j] not in "()":
                j += 1
            group = pattern[i:j]
            for symbol in group:
                suits.append([symbol])
            i = j
    return suits


def load_patterns(csv_path: Path) -> List[ParsedPattern]:
    """Load pattern rows from a CSV file."""

    with csv_path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise ValueError("CSV file is missing a header row.")

        field_map = {name.lower(): name for name in reader.fieldnames}
        try:
            pattern_field = field_map["pattern"]
            combination_field = field_map["combinations"]
        except KeyError as exc:
            raise ValueError(
                "CSV must contain 'pattern' and 'combinations' columns."
            ) from exc

        rows: List[ParsedPattern] = []
        for row in reader:
            pattern = row[pattern_field].strip()
            combinations = row[combination_field].strip()
            if not pattern:
                continue
            suit_groups = parse_pattern(pattern)
            rows.append(ParsedPattern(pattern, combinations, suit_groups))
    return rows


def build_workbook(
    parsed: Sequence[ParsedPattern],
    *,
    colour_palette: Sequence[str] | None = None,
) -> Workbook:
    """Create an Excel workbook with colour-coded cards."""

    if colour_palette is None:
        colour_palette = (
            "FFC7CE",  # light red
            "C6EFCE",  # light green
            "BDD7EE",  # light blue
            "FFE699",  # light yellow
            "F8CBAD",  # light orange
            "D9D2E9",  # lavender
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Patterns"

    max_cards = max((len(item.flattened_cards) for item in parsed), default=0)

    headers = ["pattern", "combinations"] + [f"card_{i}" for i in range(1, max_cards + 1)]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, item in enumerate(parsed, start=2):
        sheet.cell(row=row_index, column=1, value=item.pattern)
        comb_cell = sheet.cell(row=row_index, column=2, value=item.combinations)
        comb_cell.alignment = Alignment(horizontal="center")

        for col_offset, (card, suit_index) in enumerate(item.flattened_cards, start=0):
            column = 3 + col_offset
            cell = sheet.cell(row=row_index, column=column, value=card)
            colour = colour_palette[suit_index % len(colour_palette)]
            cell.fill = PatternFill(start_color=colour, end_color=colour, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    return workbook


def export_patterns(csv_path: Path, output_path: Path) -> None:
    """Load the CSV file and export the coloured Excel workbook."""

    parsed_rows = load_patterns(csv_path)
    workbook = build_workbook(parsed_rows)
    workbook.save(output_path)


def _default_output_path(input_path: Path) -> Path:
    return input_path.with_suffix(".xlsx")


def main(argv: Iterable[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Convert card patterns to a coloured Excel file.")
    parser.add_argument("csv_path", type=Path, help="Path to the input CSV file")
    parser.add_argument(
        "output_path",
        type=Path,
        nargs="?",
        help="Destination Excel file (defaults to the CSV name with .xlsx extension)",
    )

    args = parser.parse_args(argv)
    csv_path: Path = args.csv_path
    if args.output_path is None:
        output_path = _default_output_path(csv_path)
    else:
        output_path = args.output_path

    export_patterns(csv_path, output_path)


if __name__ == "__main__":  # pragma: no cover - CLI entry point
    main()
